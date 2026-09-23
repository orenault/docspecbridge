from __future__ import annotations

import json
import math
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from .canonical import new_document, validate_document
from .i18n import tr
from .package_io import write_canonical_package
from .utils import safe_stem, write_json

_RANGE_RE = re.compile(r"^(?:'(?P<quoted>[^']+)'|(?P<plain>[^!]+))!(?P<range>\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?)$")


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def _cell_blocks(value: str) -> list[dict[str, Any]]:
    return [{"type": "paragraph", "inlines": [{"type": "text", "text": value, "marks": []}]}]


def _formula_ref_values(workbook, formula: str | None) -> list[Any]:
    if not formula:
        return []
    match = _RANGE_RE.match(str(formula).replace("$", ""))
    if not match:
        return []
    sheet_name = match.group("quoted") or match.group("plain")
    if sheet_name not in workbook.sheetnames:
        return []
    ws = workbook[sheet_name]
    min_col, min_row, max_col, max_row = range_boundaries(match.group("range").replace("$", ""))
    values: list[Any] = []
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            values.append(cell.value)
    return values


def _chart_title(chart: Any, fallback: str) -> str:
    try:
        tx = chart.title.tx.rich.p[0].r[0].t
        if tx:
            return str(tx)
    except Exception:
        pass
    return fallback


def _series_refs(series: Any) -> tuple[str | None, str | None]:
    values = None
    categories = None
    try:
        values = series.val.numRef.f
    except Exception:
        try:
            values = series.yVal.numRef.f
        except Exception:
            pass
    try:
        categories = series.cat.strRef.f
    except Exception:
        try:
            categories = series.cat.numRef.f
        except Exception:
            try:
                categories = series.xVal.numRef.f
            except Exception:
                pass
    return categories, values


def _svg_escape(value: Any) -> str:
    import html
    return html.escape(_text(value), quote=True)


def _render_chart_svg(chart: Any, workbook_values, path: Path) -> dict[str, Any]:
    """Render common Excel chart types to a portable SVG without an office runtime.

    The original XLSX remains authoritative. Unsupported charts still get metadata in
    xlsx.workbook.json; common Bar/Line/Pie charts get a lightweight visual preview.
    """
    kind = chart.__class__.__name__.lower()
    title = _chart_title(chart, chart.__class__.__name__)
    series_rows: list[dict[str, Any]] = []
    for index, series in enumerate(getattr(chart, "ser", []) or [], 1):
        cat_ref, val_ref = _series_refs(series)
        cats = _formula_ref_values(workbook_values, cat_ref)
        vals = _formula_ref_values(workbook_values, val_ref)
        numeric = []
        for value in vals:
            try:
                numeric.append(float(value) if value is not None else 0.0)
            except (TypeError, ValueError):
                numeric.append(0.0)
        series_rows.append({"name": f"Series {index}", "categories": cats, "values": numeric, "category_ref": cat_ref, "value_ref": val_ref})

    meta = {"type": chart.__class__.__name__, "title": title, "series": series_rows, "preview": None}
    if not series_rows or not any(row["values"] for row in series_rows):
        return meta

    width, height = 720, 360
    left, top, right, bottom = 58, 44, 24, 54
    plot_w, plot_h = width-left-right, height-top-bottom
    palette = ["#4E79A7", "#F28E2B", "#E15759", "#76B7B2", "#59A14F", "#EDC948"]
    body: list[str] = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">', '<rect width="100%" height="100%" fill="white"/>', f'<text x="{width/2}" y="25" text-anchor="middle" font-family="Arial" font-size="18">{_svg_escape(title)}</text>']

    if "pie" in kind:
        vals = series_rows[0]["values"]
        cats = series_rows[0]["categories"] or list(range(1, len(vals)+1))
        total = sum(max(0.0, v) for v in vals) or 1.0
        cx, cy, radius = 225, 190, 120
        angle = -math.pi / 2
        for idx, value in enumerate(vals):
            portion = max(0.0, value) / total
            next_angle = angle + portion * 2 * math.pi
            x1, y1 = cx + radius * math.cos(angle), cy + radius * math.sin(angle)
            x2, y2 = cx + radius * math.cos(next_angle), cy + radius * math.sin(next_angle)
            large = 1 if portion > 0.5 else 0
            body.append(f'<path d="M {cx} {cy} L {x1:.2f} {y1:.2f} A {radius} {radius} 0 {large} 1 {x2:.2f} {y2:.2f} Z" fill="{palette[idx % len(palette)]}"/>')
            label = cats[idx] if idx < len(cats) else idx + 1
            body.append(f'<rect x="470" y="{75 + idx*24}" width="14" height="14" fill="{palette[idx % len(palette)]}"/><text x="492" y="{87 + idx*24}" font-family="Arial" font-size="12">{_svg_escape(label)} ({value:g})</text>')
            angle = next_angle
    else:
        max_value = max((max(row["values"] or [0]) for row in series_rows), default=1.0) or 1.0
        body.extend([f'<line x1="{left}" y1="{top}" x2="{left}" y2="{top+plot_h}" stroke="#555"/>', f'<line x1="{left}" y1="{top+plot_h}" x2="{left+plot_w}" y2="{top+plot_h}" stroke="#555"/>'])
        n = max(len(row["values"]) for row in series_rows)
        if "bar" in kind:
            group_w = plot_w / max(1, n)
            bar_w = group_w / max(1, len(series_rows)+1)
            for sidx, row in enumerate(series_rows):
                for idx, value in enumerate(row["values"]):
                    h = plot_h * value / max_value
                    x = left + idx*group_w + (sidx+0.5)*bar_w
                    y = top + plot_h - h
                    body.append(f'<rect x="{x:.2f}" y="{y:.2f}" width="{bar_w*0.85:.2f}" height="{h:.2f}" fill="{palette[sidx % len(palette)]}"/>')
        else:  # line/scatter-like preview
            for sidx, row in enumerate(series_rows):
                points = []
                for idx, value in enumerate(row["values"]):
                    x = left + (plot_w * idx / max(1, len(row["values"])-1))
                    y = top + plot_h - plot_h * value / max_value
                    points.append(f"{x:.2f},{y:.2f}")
                if points:
                    body.append(f'<polyline points="{" ".join(points)}" fill="none" stroke="{palette[sidx % len(palette)]}" stroke-width="3"/>')
        cats = series_rows[0]["categories"] or list(range(1, n+1))
        for idx, cat in enumerate(cats[:n]):
            x = left + (plot_w * idx / max(1, n-1))
            body.append(f'<text x="{x:.2f}" y="{height-25}" text-anchor="middle" font-family="Arial" font-size="10">{_svg_escape(cat)}</text>')

    body.append("</svg>")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(body), encoding="utf-8")
    meta["preview"] = path.name
    return meta


def extract_xlsx_to_package(path: Path, package_dir: Path, config: dict[str, Any]) -> dict[str, Any]:
    """Extract an XLSX workbook into one publication package or a workbook hierarchy.

    A one-sheet workbook is represented directly by the root package: publishing it
    creates one page containing the worksheet. A multi-sheet workbook gets a root
    index page plus one child package per worksheet. The root index is target-aware:
    human Markdown links to child Markdown, local HTML links to child HTML and the
    Confluence renderer emits the dynamic child-pages listing marker.
    """
    formula_wb = load_workbook(path, data_only=False, read_only=False)
    values_wb = load_workbook(path, data_only=True, read_only=False)
    stem = safe_stem(path.stem)
    profiles = config.get("profiles") or {}
    rag_profile = profiles.get("rag") or {}
    publication_profile = profiles.get("publication") or {}
    sheet_names = list(formula_wb.sheetnames)
    workbook_title = str(getattr(formula_wb.properties, "title", None) or "").strip() or path.stem

    workbook_meta: dict[str, Any] = {"source": path.name, "sheets": [], "defined_names": []}
    for item in formula_wb.defined_names.values():
        workbook_meta["defined_names"].append({"name": item.name, "value": item.attr_text})

    def build_sheet(
        ws,
        values_ws,
        *,
        target_dir: Path,
        output_stem: str,
        sheet_index: int,
        document_title: str | None = None,
        root_source: bool = False,
    ) -> tuple[dict[str, Path | None], dict[str, Any]]:
        source_meta = {
            "type": "xlsx" if root_source else "xlsx-sheet",
            "workbook": path.name,
            "original_path": str(path),
            "extension": ".xlsx",
            "worksheet": ws.title,
            "worksheet_index": sheet_index,
            "sheet_count": len(sheet_names),
        }
        doc = new_document(title=document_title or ws.title, source=source_meta)
        # Keep the worksheet name visible even when the single-page Confluence title
        # comes from the workbook/document title.
        doc["blocks"].append({
            "type": "heading", "level": 1,
            "inlines": [{"type": "text", "text": ws.title, "marks": []}],
        })

        formula_meta: list[dict[str, Any]] = []
        rows: list[dict[str, Any]] = []
        max_row = ws.max_row or 0
        max_col = ws.max_column or 0
        for r in range(1, max_row + 1):
            cells: list[dict[str, Any]] = []
            for c in range(1, max_col + 1):
                cell = ws.cell(r, c)
                cached = values_ws.cell(r, c).value
                display = cached if cell.data_type == "f" and cached is not None else cell.value
                if cell.data_type == "f":
                    formula_meta.append({
                        "cell": cell.coordinate,
                        "formula": cell.value,
                        "cached_value": cached,
                        "cached_value_missing": cached is None,
                    })
                    if cached is None:
                        display = tr(config, "xlsx.cached_unavailable", formula=cell.value)
                cells.append({
                    "type": "table_cell", "header": r == 1, "colspan": 1, "rowspan": 1, "style": {},
                    "blocks": _cell_blocks(_text(display)),
                })
            rows.append({"type": "table_row", "cells": cells})
        if rows and max_col:
            doc["blocks"].append({"type": "table", "rows": rows, "logical_columns": max_col})
        else:
            doc["blocks"].append({
                "type": "paragraph",
                "inlines": [{"type": "text", "text": tr(config, "xlsx.empty_worksheet"), "marks": []}],
            })

        chart_meta: list[dict[str, Any]] = []
        for chart_index, chart in enumerate(getattr(ws, "_charts", []) or [], 1):
            svg = target_dir / "images" / f"chart_{chart_index:03d}.svg"
            meta = _render_chart_svg(chart, values_wb, svg)
            chart_meta.append(meta)
            if meta.get("preview"):
                rel = str(Path("images") / svg.name).replace("\\", "/")
                doc["assets"].append({"role": "chart", "file": rel, "saved": True, "mime_type": "image/svg+xml"})
                doc["blocks"].append({
                    "type": "image", "src": rel,
                    "alt": str(meta.get("title") or f"Chart {chart_index}"), "role": "chart",
                })

        validate_document(doc)
        outputs = write_canonical_package(
            doc, target_dir, stem=output_stem,
            rag_profile=rag_profile, publication_profile=publication_profile,
        )
        sheet_meta = {
            "name": ws.title,
            "index": sheet_index,
            "max_row": max_row,
            "max_column": max_col,
            "formulas": formula_meta,
            "charts": chart_meta,
        }
        return outputs, sheet_meta

    # One worksheet: the root package *is* the worksheet page. There is no artificial
    # workbook landing page and therefore no child hierarchy to publish.
    if len(sheet_names) == 1:
        ws = formula_wb.worksheets[0]
        outputs, sheet_meta = build_sheet(
            ws, values_wb[ws.title], target_dir=package_dir, output_stem=stem,
            sheet_index=1, document_title=workbook_title, root_source=True,
        )
        sheet_meta["package"] = "."
        workbook_meta["sheets"].append(sheet_meta)
        write_json(package_dir / f"{stem}.workbook.json", workbook_meta)
        manifest_path = package_dir / "manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest["xlsx"] = {
            "workbook_metadata": f"{stem}.workbook.json",
            "single_sheet": True,
            "worksheet": ws.title,
            "children": [],
        }
        manifest["children"] = []
        write_json(manifest_path, manifest)
        return {"root": outputs, "children": [], "metadata": workbook_meta}

    # Multi-sheet workbook: create each worksheet package first so the landing page can
    # contain real relative links to the generated local Markdown/HTML artifacts.
    children: list[dict[str, Any]] = []
    for sheet_index, ws in enumerate(formula_wb.worksheets, 1):
        values_ws = values_wb[ws.title]
        sheet_stem = safe_stem(ws.title) or f"sheet-{sheet_index}"
        child_dir = package_dir / "sheets" / f"{sheet_index:02d}-{sheet_stem}"
        child_dir.mkdir(parents=True, exist_ok=True)
        outputs, sheet_meta = build_sheet(
            ws, values_ws, target_dir=child_dir, output_stem=sheet_stem,
            sheet_index=sheet_index,
        )
        package_rel = str(child_dir.relative_to(package_dir)).replace("\\", "/")
        sheet_meta["package"] = package_rel
        workbook_meta["sheets"].append(sheet_meta)
        children.append({
            "title": ws.title,
            "package": package_rel,
            "manifest": str((child_dir / "manifest.json").relative_to(package_dir)).replace("\\", "/"),
            "human_markdown": str(outputs["human_markdown"].relative_to(package_dir)).replace("\\", "/"),
            "html": str(outputs["html"].relative_to(package_dir)).replace("\\", "/"),
            "publication_markdown": str(outputs["confluence_markdown"].relative_to(package_dir)).replace("\\", "/"),
        })

    root = new_document(
        title=workbook_title,
        source={"type": "xlsx", "original_path": str(path), "extension": ".xlsx", "sheet_count": len(sheet_names)},
    )
    root["blocks"] = [
        {"type": "heading", "level": 1, "inlines": [{"type": "text", "text": workbook_title, "marks": []}]},
        {"type": "paragraph", "inlines": [{
            "type": "text",
            "text": tr(config, "xlsx.workbook_summary", count=len(sheet_names)),
            "marks": [],
        }]},
        {"type": "heading", "level": 2, "inlines": [{"type": "text", "text": tr(config, "xlsx.contents"), "marks": []}]},
        {"type": "child_pages", "items": [
            {
                "title": child["title"],
                "markdown_href": child["human_markdown"],
                "html_href": child["html"],
                "package": child["package"],
            }
            for child in children
        ]},
    ]
    root_outputs = write_canonical_package(
        root, package_dir, stem=stem,
        rag_profile=rag_profile, publication_profile=publication_profile,
    )

    write_json(package_dir / f"{stem}.workbook.json", workbook_meta)
    manifest_path = package_dir / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["xlsx"] = {
        "workbook_metadata": f"{stem}.workbook.json",
        "single_sheet": False,
        "children": children,
    }
    manifest["children"] = children
    write_json(manifest_path, manifest)
    return {"root": root_outputs, "children": children, "metadata": workbook_meta}
