from __future__ import annotations

import hashlib
import io
import math
import posixpath
import re
import zipfile
from collections import defaultdict
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from PIL import Image

EMU_PER_INCH = 914400
CSS_PX_PER_INCH = 96
PT_PER_INCH = 72

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "v": "urn:schemas-microsoft-com:vml",
}


def _sha(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def _normalize_palette_transparency(image: Image.Image) -> Image.Image:
    """Convert palette images with transparency before Pillow operations.

    Pillow warns when palette (P) images carry transparency as palette metadata /
    bytes and are converted or resized directly.  RGBA preserves the visible result
    and avoids the warning.
    """
    if image.mode == "P" and "transparency" in image.info:
        return image.convert("RGBA")
    return image


def _dhash_bytes(payload: bytes) -> str | None:
    try:
        with Image.open(io.BytesIO(payload)) as image:
            image = _normalize_palette_transparency(image)
            image = image.convert("L").resize((9, 8))
            pixels = list(image.getdata())
        bits = []
        for y in range(8):
            row = pixels[y * 9 : (y + 1) * 9]
            bits.extend(row[x] > row[x + 1] for x in range(8))
        value = 0
        for bit in bits:
            value = (value << 1) | int(bit)
        return f"{value:016x}"
    except Exception:
        return None


def _image_dimensions(path: Path) -> tuple[int | None, int | None]:
    try:
        with Image.open(path) as image:
            return int(image.width), int(image.height)
    except Exception:
        return None, None


def enrich_asset_native_dimensions(asset_meta: list[dict[str, Any]], package_dir: Path) -> None:
    seen: set[str] = set()
    for item in asset_meta:
        rel = str(item.get("file") or "")
        if not rel or rel in seen:
            continue
        seen.add(rel)
        path = package_dir / rel
        if not path.is_file():
            continue
        width, height = _image_dimensions(path)
        if width and height:
            item.setdefault("native", {})["width_px"] = width
            item.setdefault("native", {})["height_px"] = height


class AssetMatcher:
    def __init__(self, asset_meta: list[dict[str, Any]], package_dir: Path):
        self.package_dir = package_dir
        self.by_sha: dict[str, str] = {}
        self.by_dhash: dict[str, list[str]] = defaultdict(list)
        for item in asset_meta:
            rel = str(item.get("file") or "")
            if not rel:
                continue
            path = package_dir / rel
            if not path.is_file():
                continue
            payload = path.read_bytes()
            digest = str(item.get("sha256") or _sha(payload))
            self.by_sha[digest] = rel
            dhash = _dhash_bytes(payload)
            if dhash:
                self.by_dhash[dhash].append(rel)

    def match(self, payload: bytes) -> str | None:
        direct = self.by_sha.get(_sha(payload))
        if direct:
            return direct
        dhash = _dhash_bytes(payload)
        if dhash:
            candidates = self.by_dhash.get(dhash) or []
            if len(candidates) == 1:
                return candidates[0]
        return None


def _emu_to_px(value: int | float) -> int:
    return max(1, int(round(float(value) / EMU_PER_INCH * CSS_PX_PER_INCH)))


def _pt_to_px(value: int | float) -> int:
    return max(1, int(round(float(value) / PT_PER_INCH * CSS_PX_PER_INCH)))


def _pptx_occurrences(source: Path, matcher: AssetMatcher) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    from pptx import Presentation

    occurrences: list[dict[str, Any]] = []
    stats = {"pictures": 0, "connectors": 0, "groups": 0, "shapes": 0, "charts": 0}
    presentation = Presentation(str(source))

    def walk(shapes: Any, slide_no: int, group_path: str = "") -> None:
        for shape in shapes:
            stats["shapes"] += 1
            shape_type_name = str(getattr(getattr(shape, "shape_type", None), "name", ""))
            if "GROUP" in shape_type_name.upper():
                stats["groups"] += 1
                child_shapes = getattr(shape, "shapes", None)
                if child_shapes is not None:
                    walk(child_shapes, slide_no, f"{group_path}/{getattr(shape, 'shape_id', '')}")
                continue
            if "CONNECTOR" in shape_type_name.upper():
                stats["connectors"] += 1
            if bool(getattr(shape, "has_chart", False)):
                stats["charts"] += 1
            try:
                image = shape.image
                payload = image.blob
            except Exception:
                continue
            stats["pictures"] += 1
            asset = matcher.match(payload)
            if not asset:
                continue
            occurrences.append(
                {
                    "asset": asset,
                    "role": "body",
                    "source": {
                        "format": "pptx",
                        "slide": slide_no,
                        "shape_id": int(getattr(shape, "shape_id", 0) or 0),
                        "group_path": group_path or None,
                    },
                    "display": {
                        "left_px": _emu_to_px(getattr(shape, "left", 0)),
                        "top_px": _emu_to_px(getattr(shape, "top", 0)),
                        "width_px": _emu_to_px(getattr(shape, "width", 0)),
                        "height_px": _emu_to_px(getattr(shape, "height", 0)),
                        "rotation": float(getattr(shape, "rotation", 0) or 0),
                    },
                }
            )

    for slide_no, slide in enumerate(presentation.slides, start=1):
        walk(slide.shapes, slide_no)
    return occurrences, stats


def _rels_for_part(archive: zipfile.ZipFile, part_name: str) -> dict[str, str]:
    part_path = Path(part_name)
    rel_name = str(part_path.parent / "_rels" / f"{part_path.name}.rels").replace("\\", "/")
    try:
        root = ET.fromstring(archive.read(rel_name))
    except KeyError:
        return {}
    result: dict[str, str] = {}
    for rel in root:
        rid = rel.attrib.get("Id")
        target = rel.attrib.get("Target")
        if not rid or not target:
            continue
        result[rid] = posixpath.normpath(posixpath.join(posixpath.dirname(part_name), target))
    return result


def _parse_vml_dimension(style: str, name: str) -> int | None:
    match = re.search(rf"(?:^|;)\s*{re.escape(name)}\s*:\s*([0-9.]+)\s*(pt|in|px)?", style, re.I)
    if not match:
        return None
    value = float(match.group(1))
    unit = (match.group(2) or "px").lower()
    if unit == "pt":
        return _pt_to_px(value)
    if unit == "in":
        return max(1, int(round(value * CSS_PX_PER_INCH)))
    return max(1, int(round(value)))


def _docx_occurrences(source: Path, matcher: AssetMatcher) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    occurrences: list[dict[str, Any]] = []
    stats = {"drawingml_images": 0, "vml_images": 0, "anchors": 0, "inline": 0}
    with zipfile.ZipFile(source) as archive:
        names = archive.namelist()
        parts: list[tuple[str, str]] = [("word/document.xml", "body")]
        parts += [(name, "header") for name in names if re.fullmatch(r"word/header\d+\.xml", name)]
        parts += [(name, "footer") for name in names if re.fullmatch(r"word/footer\d+\.xml", name)]
        for part_name, role in parts:
            try:
                root = ET.fromstring(archive.read(part_name))
            except KeyError:
                continue
            rels = _rels_for_part(archive, part_name)
            for node in root.iter():
                if node.tag in {f"{{{NS['wp']}}}inline", f"{{{NS['wp']}}}anchor"}:
                    if node.tag.endswith("anchor"):
                        stats["anchors"] += 1
                    else:
                        stats["inline"] += 1
                    extent = node.find(f"{{{NS['wp']}}}extent")
                    blip = node.find(f".//{{{NS['a']}}}blip")
                    if blip is None:
                        continue
                    rid = blip.attrib.get(f"{{{NS['r']}}}embed")
                    media = rels.get(rid or "")
                    if not media:
                        continue
                    try:
                        payload = archive.read(media)
                    except KeyError:
                        continue
                    asset = matcher.match(payload)
                    if not asset:
                        continue
                    stats["drawingml_images"] += 1
                    width = _emu_to_px(int(extent.attrib.get("cx", "0"))) if extent is not None else None
                    height = _emu_to_px(int(extent.attrib.get("cy", "0"))) if extent is not None else None
                    occurrences.append(
                        {
                            "asset": asset,
                            "role": role,
                            "source": {"format": "docx", "part": part_name, "relationship_id": rid},
                            "display": {"width_px": width, "height_px": height},
                        }
                    )
                elif node.tag.endswith("}shape") and node.tag.startswith(f"{{{NS['v']}}}"):
                    image_data = node.find(f".//{{{NS['v']}}}imagedata")
                    if image_data is None:
                        continue
                    rid = image_data.attrib.get(f"{{{NS['r']}}}id")
                    media = rels.get(rid or "")
                    if not media:
                        continue
                    try:
                        payload = archive.read(media)
                    except KeyError:
                        continue
                    asset = matcher.match(payload)
                    if not asset:
                        continue
                    stats["vml_images"] += 1
                    style = node.attrib.get("style", "")
                    occurrences.append(
                        {
                            "asset": asset,
                            "role": role,
                            "source": {"format": "docx", "part": part_name, "relationship_id": rid, "vml": True},
                            "display": {
                                "width_px": _parse_vml_dimension(style, "width"),
                                "height_px": _parse_vml_dimension(style, "height"),
                            },
                        }
                    )
    return occurrences, stats


def _pdf_occurrences(source: Path, matcher: AssetMatcher) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    import pymupdf

    occurrences: list[dict[str, Any]] = []
    stats = {"embedded_images": 0, "vector_drawings": 0, "pages": 0}
    doc = pymupdf.open(str(source))
    try:
        stats["pages"] = len(doc)
        for page_no, page in enumerate(doc, start=1):
            try:
                stats["vector_drawings"] += len(page.get_drawings())
            except Exception:
                pass
            seen_xref: set[int] = set()
            for info in page.get_images(full=True):
                xref = int(info[0])
                if xref <= 0 or xref in seen_xref:
                    continue
                seen_xref.add(xref)
                try:
                    payload = doc.extract_image(xref)["image"]
                    rects = page.get_image_rects(xref)
                except Exception:
                    continue
                asset = matcher.match(payload)
                if not asset:
                    continue
                stats["embedded_images"] += 1
                for rect in rects:
                    occurrences.append(
                        {
                            "asset": asset,
                            "role": "body",
                            "source": {"format": "pdf", "page": page_no, "xref": xref},
                            "display": {
                                "left_px": _pt_to_px(rect.x0),
                                "top_px": _pt_to_px(rect.y0),
                                "width_px": _pt_to_px(rect.width),
                                "height_px": _pt_to_px(rect.height),
                            },
                        }
                    )
    finally:
        doc.close()
    return occurrences, stats


def collect_image_geometry(
    source: Path,
    asset_meta: list[dict[str, Any]],
    package_dir: Path,
) -> tuple[list[dict[str, Any]], dict[str, Any], list[str]]:
    """Collect source display geometry without adding layout noise to Markdown.

    Geometry is stored in manifest/document JSON. Publication can use it to create
    display-sized derivatives; RAG keeps the original high-resolution assets.
    """
    enrich_asset_native_dimensions(asset_meta, package_dir)
    matcher = AssetMatcher(asset_meta, package_dir)
    warnings: list[str] = []
    suffix = source.suffix.lower()
    try:
        if suffix == ".pptx":
            occurrences, stats = _pptx_occurrences(source, matcher)
        elif suffix == ".docx":
            occurrences, stats = _docx_occurrences(source, matcher)
        elif suffix == ".pdf":
            occurrences, stats = _pdf_occurrences(source, matcher)
        else:
            return [], {}, []
    except Exception as exc:
        return [], {}, [f"Géométrie source non analysée: {exc}"]

    by_asset: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for occ in occurrences:
        by_asset[str(occ["asset"])].append(occ)
    for item in asset_meta:
        rel = str(item.get("file") or "")
        if rel in by_asset:
            item["display_occurrences"] = by_asset[rel]

    if suffix == ".pdf" and int(stats.get("vector_drawings", 0)):
        warnings.append(
            f"PDF: {stats['vector_drawings']} tracé(s) vectoriel(s) PDF détecté(s). Le texte/images sont extraits, "
            "mais le rendu vectoriel n'est pas encore rasterisé automatiquement."
        )
    return occurrences, stats, warnings


def _valid_size(display: dict[str, Any], min_px: int, max_px: int) -> tuple[int, int] | None:
    try:
        width = int(display.get("width_px") or 0)
        height = int(display.get("height_px") or 0)
    except Exception:
        return None
    if width <= 0 or height <= 0:
        return None
    if width < min_px and height < min_px:
        scale = min_px / max(width, height)
        width = max(1, int(round(width * scale)))
        height = max(1, int(round(height * scale)))
    if max(width, height) > max_px:
        scale = max_px / max(width, height)
        width = max(1, int(round(width * scale)))
        height = max(1, int(round(height * scale)))
    return width, height


def build_publication_variants(
    markdown: str,
    package_dir: Path,
    asset_meta: list[dict[str, Any]],
    profile: dict[str, Any],
    image_regex: re.Pattern[str],
) -> tuple[str, list[dict[str, Any]], list[str]]:
    """Create source-sized raster derivatives and rewrite publication Markdown.

    Keeping the size in the pixels of a publication derivative is deliberately used
    instead of layout metadata in Markdown. This keeps the Markdown lean, gives VS Code
    a useful preview and lets md2conf upload an image whose intrinsic size is close to
    the source display size. Original assets stay untouched for RAG.
    """
    if not profile.get("preserve_image_display_size", True):
        return markdown, [], []

    out_dir_name = str(profile.get("display_image_directory") or "publication_images")
    out_dir = package_dir / out_dir_name
    out_dir.mkdir(parents=True, exist_ok=True)
    min_px = int(profile.get("min_display_px", 12))
    max_px = int(profile.get("max_display_px", 1800))
    avoid_upscale = bool(profile.get("avoid_upscale", False))

    queues: dict[str, list[dict[str, Any]]] = defaultdict(list)
    native: dict[str, tuple[int | None, int | None]] = {}
    for item in asset_meta:
        rel = str(item.get("file") or "")
        if not rel:
            continue
        native_data = item.get("native") or {}
        native[rel] = (native_data.get("width_px"), native_data.get("height_px"))
        for occ in item.get("display_occurrences") or []:
            if occ.get("role") == "body" or not queues[rel]:
                queues[rel].append(occ)

    counters: dict[str, int] = defaultdict(int)
    cache: dict[tuple[str, int, int], str] = {}
    variants: list[dict[str, Any]] = []
    warnings: list[str] = []

    def replace(match: re.Match[str]) -> str:
        target = match.group("target").strip().replace("\\", "/")
        if not target.startswith("images/"):
            return match.group(0)
        occs = queues.get(target) or []
        if not occs:
            return match.group(0)
        idx = counters[target]
        occurrence = occs[min(idx, len(occs) - 1)]
        counters[target] += 1
        size = _valid_size(occurrence.get("display") or {}, min_px, max_px)
        if not size:
            return match.group(0)
        width, height = size
        nw, nh = native.get(target, (None, None))
        if avoid_upscale and nw and nh and (width > nw or height > nh):
            width, height = int(nw), int(nh)

        source_path = package_dir / target
        if not source_path.is_file():
            return match.group(0)
        key = (target, width, height)
        rel_variant = cache.get(key)
        if rel_variant is None:
            suffix = source_path.suffix.lower()
            if suffix not in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif"}:
                return match.group(0)
            variant_name = f"{source_path.stem}__w{width}_h{height}{suffix}"
            variant_path = out_dir / variant_name
            try:
                with Image.open(source_path) as image:
                    if getattr(image, "is_animated", False):
                        image.seek(0)
                    image = _normalize_palette_transparency(image)
                    resampling = getattr(Image, "Resampling", Image).LANCZOS
                    resized = image.resize((width, height), resample=resampling)
                    if suffix in {".jpg", ".jpeg"} and resized.mode not in {"RGB", "L"}:
                        resized = resized.convert("RGB")
                    resized.save(variant_path)
            except Exception as exc:
                warnings.append(f"Taille d'affichage non appliquée à {target}: {exc}")
                return match.group(0)
            rel_variant = str(Path(out_dir_name) / variant_name).replace("\\", "/")
            cache[key] = rel_variant
            variants.append(
                {
                    "source_asset": target,
                    "file": rel_variant,
                    "width_px": width,
                    "height_px": height,
                }
            )
        alt = match.group("alt") or "Image"
        return f"![{alt}]({rel_variant})"

    result = image_regex.sub(replace, markdown)
    return result, variants, warnings
