from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

_GENERIC_TITLES = {
    "untitled",
    "document",
    "document1",
    "document 1",
    "microsoft word",
    "microsoft word document",
    "powerpoint presentation",
    "presentation",
    "presentation1",
    "presentation 1",
    "adobe acrobat",
}


def _clean(value: Any) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    if not text:
        return ""
    lowered = text.casefold()
    if lowered in _GENERIC_TITLES:
        return ""
    if lowered.startswith("microsoft word - ") and lowered.endswith((".doc", ".docx")):
        return ""
    if len(text) > 240:
        return ""
    return text


def _core_properties_title(path: Path) -> str:
    if path.suffix.lower() not in {".docx", ".pptx"}:
        return ""
    try:
        with zipfile.ZipFile(path) as archive:
            payload = archive.read("docProps/core.xml")
        root = ET.fromstring(payload)
        for elem in root.iter():
            if elem.tag.endswith("}title") or elem.tag == "title":
                title = _clean(elem.text)
                if title:
                    return title
    except Exception:
        pass
    return ""



def _docx_visual_title(path: Path) -> str:
    """Find a visually prominent early Word paragraph when core title is empty.

    Many enterprise DOCX templates leave Core Properties/title blank while rendering the
    real document title as a large bold paragraph before the first section heading. The
    heuristic is deliberately strict so a section such as "SUIVI" is not promoted.
    """
    try:
        from docx import Document

        document = Document(path)
        best: tuple[float, str] | None = None
        noisy = {
            "suivi", "historique des modifications", "documents liés",
            "table des matières", "sommaire", "contents",
        }
        for index, paragraph in enumerate(document.paragraphs[:40]):
            title = _clean(paragraph.text)
            if not title or not (5 <= len(title) <= 220):
                continue
            if title.casefold() in noisy:
                continue
            runs = [run for run in paragraph.runs if str(run.text or "").strip()]
            sizes = [float(run.font.size.pt) for run in runs if run.font.size is not None]
            max_size = max(sizes) if sizes else 0.0
            bold_count = sum(1 for run in runs if run.bold is True)
            score = 0.0
            if max_size >= 20:
                score += 6.0
            elif max_size >= 16:
                score += 4.0
            elif max_size >= 14:
                score += 2.0
            if runs and bold_count / len(runs) >= 0.6:
                score += 2.0
            if paragraph.alignment is not None and int(paragraph.alignment) in {1, 2}:
                score += 1.0
            if index < 15:
                score += 1.0
            style_name = str(getattr(getattr(paragraph, "style", None), "name", "") or "").casefold()
            if style_name in {"title", "titre"} and max_size >= 14:
                score += 1.0
            if score >= 5.0 and (best is None or score > best[0]):
                best = (score, title)
        return best[1] if best else ""
    except Exception:
        return ""

def _pptx_first_slide_title(path: Path) -> str:
    try:
        from pptx import Presentation

        presentation = Presentation(path)
        if not presentation.slides:
            return ""
        slide = presentation.slides[0]
        title_shape = getattr(slide.shapes, "title", None)
        if title_shape is not None:
            title = _clean(getattr(title_shape, "text", ""))
            if title:
                return title
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                title = _clean(getattr(shape, "text", ""))
                if title:
                    return title
    except Exception:
        pass
    return ""


def _pdf_metadata_title(path: Path) -> str:
    try:
        import fitz

        with fitz.open(path) as pdf:
            title = _clean((pdf.metadata or {}).get("title"))
            if title:
                return title
    except Exception:
        pass
    return ""


def _first_heading(doc: dict[str, Any]) -> str:
    for block in doc.get("blocks") or []:
        if block.get("type") != "heading":
            continue
        parts: list[str] = []
        for item in block.get("inlines") or []:
            if item.get("type") == "text":
                parts.append(str(item.get("text") or ""))
        title = _clean("".join(parts))
        if title:
            return title
    return ""


def detect_document_title(
    source: Path,
    doc: dict[str, Any],
    *,
    xberg_metadata: Any = None,
) -> tuple[str, str]:
    """Return a reliable document title and its source.

    The filename is the final fallback. Detection is deliberately conservative: a bad
    metadata title is worse than a predictable filename in an automated publication flow.
    """
    suffix = source.suffix.lower()

    title = _core_properties_title(source)
    if title:
        return title, "core_properties"

    if suffix == ".docx":
        title = _docx_visual_title(source)
        if title:
            return title, "document_visual_title"

    if suffix == ".pdf":
        title = _pdf_metadata_title(source)
        if title:
            return title, "pdf_metadata"

    if suffix == ".pptx":
        title = _pptx_first_slide_title(source)
        if title:
            return title, "first_slide_title"

    if isinstance(xberg_metadata, dict):
        for key in ("title", "Title", "document_title", "subject"):
            title = _clean(xberg_metadata.get(key))
            if title:
                return title, "extractor_metadata"

    title = _first_heading(doc)
    if title:
        return title, "document_heading"

    return source.stem, "filename"


def apply_detected_title(
    doc: dict[str, Any],
    source: Path,
    *,
    xberg_metadata: Any = None,
) -> tuple[str, str]:
    title, source_kind = detect_document_title(source, doc, xberg_metadata=xberg_metadata)
    doc["title"] = title
    metadata = doc.setdefault("metadata", {})
    metadata["title"] = title
    metadata["title_source"] = source_kind
    metadata["filename_title"] = source.stem
    return title, source_kind
