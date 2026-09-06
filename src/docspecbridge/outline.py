from __future__ import annotations

import difflib
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET


TOC_PLACEHOLDER = "<!-- DOCSPECBRIDGE_TOC -->"
_TOC_NAMES = {
    "sommaire",
    "table des matieres",
    "table des matières",
    "table des matieres",
    "table des matires",
    "contents",
    "table of contents",
    "inhaltsverzeichnis",
    "indice",
    "índice",
    "目录",
    "目錄",
}
_HEADING_RE = re.compile(r"^\s{0,3}(#{1,6})\s+(.+?)\s*$")
_MD_LINK_RE = re.compile(r"\[([^\]]+)\]\([^)]*\)")
_LEADING_NUMBER_RE = re.compile(r"^\s*(?:\d+(?:\.\d+)*[.)]?|[IVXLCDM]+[.)])\s+", re.I)
_TRAILING_PAGE_RE = re.compile(r"(?:\s*[.·•…]{2,}\s*|\s+)(?:\d+|[ivxlcdm]+)\s*$", re.I)
_STYLE_LEVEL_RE = re.compile(
    r"(?:heading|titre|überschrift|uberschrift|t[ií]tulo|标题|標題)[ _-]?(\d+)", re.I
)
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@dataclass
class OutlineResult:
    markdown: str
    outline: list[dict[str, Any]]
    source: str
    toc_detected: bool
    headings_applied: int
    warnings: list[str]


def _clean_inline_markdown(value: str) -> str:
    value = value.strip()
    heading = _HEADING_RE.match(value)
    if heading:
        value = heading.group(2)
    value = _MD_LINK_RE.sub(r"\1", value)
    value = re.sub(r"[`*_~]", "", value)
    value = _LEADING_NUMBER_RE.sub("", value)
    value = _TRAILING_PAGE_RE.sub("", value)
    value = re.sub(r"\s+", " ", value).strip(" .:-–—\t")
    return value


def _norm(value: str) -> str:
    value = _clean_inline_markdown(value).casefold()
    value = value.replace("’", "'").replace("–", "-").replace("—", "-")
    value = "".join(ch for ch in unicodedata.normalize("NFKD", value) if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", value).strip()


def _outline_from_markdown(markdown: str) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for line_no, line in enumerate(markdown.splitlines(), 1):
        match = _HEADING_RE.match(line)
        if not match:
            continue
        result.append(
            {
                "level": len(match.group(1)),
                "title": _clean_inline_markdown(match.group(2)),
                "line": line_no,
                "source": "markdown",
                "confidence": 0.8,
            }
        )
    return result


def _docx_style_levels(archive: zipfile.ZipFile) -> dict[str, int]:
    result: dict[str, int] = {}
    try:
        root = ET.fromstring(archive.read("word/styles.xml"))
    except Exception:
        return result
    for style in root.findall(f".//{_W}style"):
        style_id = style.attrib.get(f"{_W}styleId", "")
        if not style_id:
            continue
        level: int | None = None
        outline = style.find(f"./{_W}pPr/{_W}outlineLvl")
        if outline is not None:
            try:
                level = int(outline.attrib.get(f"{_W}val", "0")) + 1
            except ValueError:
                pass
        name_node = style.find(f"./{_W}name")
        name = name_node.attrib.get(f"{_W}val", "") if name_node is not None else ""
        if level is None:
            match = _STYLE_LEVEL_RE.search(f"{style_id} {name}")
            if match:
                level = int(match.group(1))
        if level is not None and 1 <= level <= 9:
            result[style_id] = level
    return result


def _extract_docx_outline(source: Path) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    try:
        with zipfile.ZipFile(source) as archive:
            style_levels = _docx_style_levels(archive)
            root = ET.fromstring(archive.read("word/document.xml"))
    except Exception:
        return result

    for paragraph in root.findall(f".//{_W}p"):
        text = "".join((node.text or "") for node in paragraph.findall(f".//{_W}t")).strip()
        if not text:
            continue
        level: int | None = None
        ppr = paragraph.find(f"./{_W}pPr")
        if ppr is not None:
            outline = ppr.find(f"./{_W}outlineLvl")
            if outline is not None:
                try:
                    level = int(outline.attrib.get(f"{_W}val", "0")) + 1
                except ValueError:
                    pass
            if level is None:
                style = ppr.find(f"./{_W}pStyle")
                style_id = style.attrib.get(f"{_W}val", "") if style is not None else ""
                level = style_levels.get(style_id)
                if level is None and style_id:
                    match = _STYLE_LEVEL_RE.search(style_id)
                    if match:
                        level = int(match.group(1))
        if level is not None and 1 <= level <= 6:
            result.append(
                {
                    "level": level,
                    "title": text,
                    "source": "docx_outline",
                    "confidence": 1.0,
                }
            )
    return result


def _extract_pdf_outline(source: Path) -> list[dict[str, Any]]:
    try:
        import pymupdf

        with pymupdf.open(source) as document:
            toc = document.get_toc(simple=True)
    except Exception:
        return []
    result: list[dict[str, Any]] = []
    for row in toc or []:
        if len(row) < 3:
            continue
        level, title, page = row[:3]
        try:
            level = max(1, min(6, int(level)))
        except Exception:
            continue
        result.append(
            {
                "level": level,
                "title": str(title).strip(),
                "page": int(page) if str(page).isdigit() else page,
                "source": "pdf_outline",
                "confidence": 1.0,
            }
        )
    return [item for item in result if item["title"]]



def _extract_pdf_link_outline(source: Path) -> list[dict[str, Any]]:
    """Infer an outline from a PDF table of contents made of internal links.

    Many Office-generated PDFs have no PDF bookmark outline, but their visible table of
    contents contains GoTo links.  Link indentation is a strong level signal and the
    destination page/position gives useful provenance.
    """
    try:
        import pymupdf

        document = pymupdf.open(source)
    except Exception:
        return []
    candidates: list[dict[str, Any]] = []
    try:
        for page_no, page in enumerate(document, 1):
            internal = [link for link in page.get_links() if int(link.get("kind", 0)) == 1 and "page" in link and "from" in link]
            if len(internal) < 3:
                continue
            for link in internal:
                rect = link.get("from")
                try:
                    text = page.get_textbox(rect).strip()
                except Exception:
                    text = ""
                toc_title = _clean_inline_markdown(text)
                if not toc_title or _norm(toc_title) in {_norm(item) for item in _TOC_NAMES}:
                    continue
                target_page_index = int(link.get("page", 0))
                target_y = float(getattr(link.get("to"), "y", 0.0))
                target_title = ""
                if 0 <= target_page_index < document.page_count:
                    target_page = document[target_page_index]
                    box = pymupdf.Rect(0, max(0, target_y - 8), target_page.rect.width, min(target_page.rect.height, target_y + 85))
                    target_lines = [line.strip() for line in target_page.get_textbox(box).splitlines() if line.strip()]
                    ranked: list[tuple[float, str]] = []
                    for target_line in target_lines:
                        if re.fullmatch(r"(?:\d+(?:\.\d+)*[.)]?|[IVXLCDM]+[.)]?)", target_line, re.I):
                            continue
                        candidate_title = _clean_inline_markdown(target_line)
                        if not candidate_title:
                            continue
                        score = difflib.SequenceMatcher(None, _norm(toc_title), _norm(candidate_title)).ratio()
                        ranked.append((score, candidate_title))
                    if ranked:
                        ranked.sort(key=lambda item: item[0], reverse=True)
                        target_title = ranked[0][1] if ranked[0][0] >= 0.40 else ranked[-1][1]
                candidates.append({
                    "title": target_title or toc_title,
                    "toc_title": toc_title,
                    "source_page": page_no,
                    "x0": round(float(rect.x0), 1),
                    "page": target_page_index + 1,
                    "target_y": round(target_y, 1),
                })
    finally:
        document.close()
    if not candidates:
        return []
    # Map indentation bands to heading levels. Office-generated TOCs typically use
    # consistent 12pt indentation, but rounding/unique ordering avoids hard-coding it.
    x_values = sorted({item["x0"] for item in candidates})
    x_level = {x: min(6, idx + 1) for idx, x in enumerate(x_values)}
    result: list[dict[str, Any]] = []
    for item in candidates:
        result.append({
            "level": x_level[item["x0"]],
            "title": item["title"],
            "toc_title": item.get("toc_title"),
            "page": item["page"],
            "target_y": item["target_y"],
            "toc_page": item["source_page"],
            "source": "pdf_internal_toc",
            "confidence": 0.95,
        })
    return result

def _extract_pptx_outline(source: Path) -> list[dict[str, Any]]:
    try:
        from pptx import Presentation

        presentation = Presentation(str(source))
    except Exception:
        return []
    result: list[dict[str, Any]] = []
    for slide_no, slide in enumerate(presentation.slides, 1):
        title_shape = slide.shapes.title
        text = title_shape.text.strip() if title_shape is not None and getattr(title_shape, "text", None) else ""
        if text:
            result.append(
                {
                    "level": 1,
                    "title": text,
                    "slide": slide_no,
                    "source": "pptx_slide_title",
                    "confidence": 1.0,
                }
            )
    return result


def extract_outline(source: Path, markdown: str) -> tuple[list[dict[str, Any]], str]:
    suffix = source.suffix.lower()
    if suffix == ".docx":
        outline = _extract_docx_outline(source)
        if outline:
            return outline, "docx_outline"
    elif suffix == ".pdf":
        outline = _extract_pdf_outline(source)
        if outline:
            return outline, "pdf_outline"
        outline = _extract_pdf_link_outline(source)
        if outline:
            return outline, "pdf_internal_toc"
    elif suffix == ".pptx":
        outline = _extract_pptx_outline(source)
        if outline:
            return outline, "pptx_slide_titles"
    markdown_outline = _outline_from_markdown(markdown)
    return markdown_outline, "markdown" if markdown_outline else "none"


def _is_toc_heading(line: str) -> bool:
    return _norm(line) in {_norm(item) for item in _TOC_NAMES}


def _find_title_matches(lines: list[str], title: str, start: int = 0) -> list[int]:
    wanted = _norm(title)
    if not wanted:
        return []
    matches: list[int] = []
    for idx in range(start, len(lines)):
        value = _norm(lines[idx])
        if value == wanted:
            matches.append(idx)
    if matches:
        return matches

    # PDF text extraction may lose accents or a few glyphs.  For reliable source outlines,
    # permit a conservative fuzzy match against short standalone lines, never paragraphs
    # or table rows.  Sequential matching in apply_outline further limits false positives.
    fuzzy: list[tuple[float, int]] = []
    for idx in range(start, len(lines)):
        raw = lines[idx].strip()
        if not raw or raw.startswith("|") or len(raw) > max(180, len(title) * 2):
            continue
        value = _norm(raw)
        if not value:
            continue
        score = difflib.SequenceMatcher(None, wanted, value).ratio()
        if score >= 0.84:
            fuzzy.append((score, idx))
    if not fuzzy:
        return []
    best = max(score for score, _ in fuzzy)
    return [idx for score, idx in fuzzy if score >= best - 0.02]


def apply_outline(markdown: str, outline: list[dict[str, Any]], *, replace_source_toc: bool = True) -> OutlineResult:
    """Apply reliable source outline levels to Markdown and isolate a source TOC.

    The original source file remains in the package, while the semantic Markdown is
    normalized for publication/RAG.  When a source TOC is detected, it is replaced by a
    neutral placeholder. Publication turns the placeholder into md2conf's ``[[_TOC_]]``;
    the RAG profile removes it entirely to avoid duplicate retrieval text.
    """
    if not outline:
        return OutlineResult(markdown.rstrip() + "\n", [], "none", False, 0, [])

    lines = markdown.splitlines()
    warnings: list[str] = []
    toc_idx = next((idx for idx, line in enumerate(lines[: max(30, len(lines) // 3)]) if _is_toc_heading(line)), None)
    toc_detected = False
    search_start = 0

    if toc_idx is not None and replace_source_toc and outline:
        body_idx: int | None = None
        # The first outline entries may precede the visible TOC (cover/revision sections).
        # Find any outline title that occurs once in the TOC area and again in the body.
        for item in outline:
            titles = [str(item.get("toc_title") or ""), str(item.get("title") or "")]
            for title in titles:
                if not title:
                    continue
                occurrences = _find_title_matches(lines, title, toc_idx + 1)
                if len(occurrences) >= 2:
                    candidate = occurrences[-1]
                    if candidate > toc_idx:
                        body_idx = candidate if body_idx is None else min(body_idx, candidate)
            if body_idx is not None:
                break
        if body_idx is not None and body_idx > toc_idx:
            lines[toc_idx:body_idx] = [TOC_PLACEHOLDER, ""]
            toc_detected = True

    applied = 0
    cursor = 0
    for item in outline:
        title = str(item.get("title") or "").strip()
        if not title:
            continue
        level = max(1, min(6, int(item.get("level") or 1)))
        matches = _find_title_matches(lines, title, cursor)
        if not matches:
            continue
        idx = matches[0]
        clean_title = _clean_inline_markdown(lines[idx]) or title
        lines[idx] = f"{'#' * level} {clean_title}"
        cursor = idx + 1
        applied += 1

    if applied == 0 and outline:
        warnings.append("Structure de titres détectée dans la source mais aucun titre n'a pu être recollé au Markdown.")

    enriched = [dict(item) for item in outline]
    for item in enriched:
        item.setdefault("confidence", 1.0 if item.get("source") != "markdown" else 0.8)
    return OutlineResult("\n".join(lines).rstrip() + "\n", enriched, str(outline[0].get("source") or "outline"), toc_detected, applied, warnings)


def publication_toc(markdown: str, enabled: str | bool = "auto") -> str:
    if TOC_PLACEHOLDER not in markdown:
        return markdown
    if enabled is False or str(enabled).lower() in {"false", "off", "no", "0"}:
        return markdown.replace(TOC_PLACEHOLDER, "")
    # auto/true: source TOC existed, therefore a Confluence-native TOC is appropriate.
    return markdown.replace(TOC_PLACEHOLDER, "[[_TOC_]]")


def rag_without_toc(markdown: str) -> str:
    return markdown.replace(TOC_PLACEHOLDER, "")
