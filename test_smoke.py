from pathlib import Path

from docspecbridge.config import get_confluence_instance, load_config
from docspecbridge.ooxml import inspect_ooxml
from docspecbridge.rag import build_rag_markdown, chunk_markdown


def test_default_config():
    cfg = load_config(None)
    assert "extensions" not in cfg["app"]
    from docspecbridge.config import SUPPORTED_SOURCE_EXTENSIONS
    assert ".docx" in SUPPORTED_SOURCE_EXTENSIONS
    assert ".xlsx" in SUPPORTED_SOURCE_EXTENSIONS
    assert cfg["extract"]["images"]["extract_images"] is True
    assert cfg["profiles"]["publication"]["preserve_image_display_size"] is True
    assert cfg["profiles"]["rag"]["enabled"] is True
    assert cfg["profiles"]["publication"]["table_of_contents"]["enabled"] == "auto"
    assert cfg["rag_export"]["destination"] == "./rag"


def test_non_ooxml():
    assert inspect_ooxml(Path("x.pdf")) is None


def test_rag_removes_header_only_image():
    md = "![](images/logo.png)\n\n# Title\n\nText\n\n![](images/figure.png)\n"
    assets = [
        {"file": "images/logo.png", "role": "header"},
        {"file": "images/figure.png", "role": "body"},
    ]
    result = build_rag_markdown(md, assets, {"include_header_images": False, "keep_image_references": True})
    assert "logo.png" not in result
    assert "figure.png" in result


def test_heading_aware_chunks():
    chunks = chunk_markdown("# A\n\nalpha\n\n## B\n\nbeta", max_characters=300, overlap=0)
    assert chunks
    assert chunks[-1]["heading_path"] == ["A", "B"]
    assert chunks[-1]["embedding_text"].startswith("A > B")


def test_multiple_confluence_instances():
    cfg = load_config(None)
    cfg["confluence"]["instances"] = {
        "prod": {"domain": "prod.atlassian.net", "token_env": "PROD_PAT"},
        "sandbox": {"domain": "sandbox.atlassian.net", "token_env": "SB_PAT"},
    }
    cfg["confluence"]["default_instance"] = "prod"
    name, instance = get_confluence_instance(cfg)
    assert name == "prod"
    assert instance["domain"] == "prod.atlassian.net"


def test_domain_normalization():
    from docspecbridge.config import normalize_domain

    assert normalize_domain("https://company.atlassian.net/") == "company.atlassian.net"
    assert normalize_domain("company.atlassian.net") == "company.atlassian.net"


def test_scoped_instance_url_migration():
    cfg = load_config(None)
    cfg["confluence"]["instances"] = {
        "prod": {
            "domain": "https://company.atlassian.net/",
            "auth_type": "scoped",
            "cloud_id": "abc-123",
            "user_name": "user@example.com",
        }
    }
    cfg["confluence"]["default_instance"] = "prod"
    _, instance = get_confluence_instance(cfg)
    assert instance["domain"] == "company.atlassian.net"
    assert instance["api_url"] == "https://api.atlassian.com/ex/confluence/abc-123"


def test_outline_replaces_source_toc_and_applies_levels():
    from docspecbridge.outline import apply_outline, publication_toc, rag_without_toc

    md = "Sommaire\n\n1. Introduction ..... 2\n1.1 Scope ..... 3\n\nIntroduction\n\nBody\n\nScope\n\nText\n"
    outline = [
        {"level": 1, "title": "Introduction", "source": "pdf_outline"},
        {"level": 2, "title": "Scope", "source": "pdf_outline"},
    ]
    result = apply_outline(md, outline)
    assert result.toc_detected is True
    assert "[[_TOC_]]" in publication_toc(result.markdown)
    assert "# Introduction" in result.markdown
    assert "## Scope" in result.markdown
    assert "Sommaire" not in rag_without_toc(result.markdown)


def test_rag_export_builds_global_corpus(tmp_path):
    import json
    from docspecbridge.rag_export import export_rag_corpus

    package = tmp_path / "out" / "doc__pdf"
    package.mkdir(parents=True)
    (package / "doc.rag.md").write_text("# Title\n\nText\n", encoding="utf-8")
    (package / "document.json").write_text("{}\n", encoding="utf-8")
    (package / "chunks.jsonl").write_text(json.dumps({"id": 1, "heading_path": ["Title"], "content": "Text", "source": {"file": "doc.pdf"}}) + "\n", encoding="utf-8")
    (package / "manifest.json").write_text(json.dumps({
        "source": {"sha256": "a" * 64, "original_path": "doc.pdf"},
        "outputs": {"rag_markdown": "doc.rag.md", "chunks": "chunks.jsonl", "document_json": "document.json"},
        "assets": {"images_directory": "images"},
        "outline": {"entries": []},
    }), encoding="utf-8")
    result = export_rag_corpus(tmp_path / "out", tmp_path / "rag")
    assert result.document_count == 1
    assert result.chunk_count == 1
    assert (tmp_path / "rag" / "corpus.json").is_file()
    assert (tmp_path / "rag" / "chunks.jsonl").is_file()


def test_palette_transparency_dhash_has_no_pillow_warning():
    import io
    import warnings
    from PIL import Image
    from docspecbridge.geometry import _dhash_bytes

    image = Image.new("P", (4, 4), 0)
    palette = [0, 0, 0, 255, 255, 255] + [0, 0, 0] * 254
    image.putpalette(palette)
    image.info["transparency"] = bytes([0, 255] + [255] * 254)
    payload = io.BytesIO()
    image.save(payload, format="PNG")

    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        assert _dhash_bytes(payload.getvalue()) is not None
    assert not [w for w in caught if "Palette images with Transparency" in str(w.message)]


def test_docx_publication_defaults():
    cfg = load_config(None)
    docx = cfg["profiles"]["publication"]["docx"]
    assert docx["engine"] == "mammoth-html"
    assert docx["preserve_merged_cells"] is True
    assert docx["preserve_lists"] is True
    assert docx["preserve_highlight_colors"] is True
    assert docx["preserve_cell_shading"] is True
    assert docx["preserve_cell_text_color"] is True


def test_md2conf_converter_defaults_are_explicit_and_safe():
    cfg = load_config(None)
    conv = cfg["confluence"]["converter"]
    assert conv["force_valid_url"] is True
    assert conv["prefer_raster"] is True
    assert conv["render_drawio"] is False
    assert conv["render_mermaid"] is False
    assert conv["render_plantuml"] is False
    assert conv["render_latex"] is False
    assert conv["diagram_output_format"] == "png"
    assert conv["user_mentions"] is True
    assert conv["force_valid_language"] is True


def test_docx_table_style_extraction_and_merged_cell_alignment(tmp_path):
    from docx import Document
    from docx.dml.color import RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docspecbridge.docx_publication import _docx_table_cell_styles

    path = tmp_path / "merged.docx"
    document = Document()
    table = document.add_table(rows=2, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "Merged"
    tc_pr = merged._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "808080")
    tc_pr.append(shd)
    run = merged.paragraphs[0].runs[0]
    run.font.color.rgb = RGBColor(255, 255, 255)
    table.cell(0, 2).text = "Other"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "B"
    table.cell(1, 2).text = "C"
    document.save(path)

    styles = _docx_table_cell_styles(path)
    assert len(styles) == 1
    assert len(styles[0][0]) == 2  # merged cell + third cell
    assert styles[0][0][0].background == "#808080"
    assert styles[0][0][0].foreground == "#FFFFFF"


def test_docx_publication_normalizes_checkboxes_unsafe_links_and_xml_tags():
    from lxml import html
    from docspecbridge.docx_publication import _normalize_mammoth_html, _serialize_for_md2conf

    root = html.fragment_fromstring(
        '<p><input type="checkbox" checked="checked"/> Done '
        '<input type="checkbox"/> Todo</p>'
        '<p><a href="javascript:alert(1)">bad</a></p>'
        '<p><img src="images/a.png" width="24" height="24"></p>',
        create_parent="div",
    )
    stats = _normalize_mammoth_html(root)
    rendered = _serialize_for_md2conf(root)
    assert stats["checkboxes"] == 2
    assert stats["unsafe_links_removed"] == 1
    assert "☒" in rendered and "☐" in rendered
    assert "javascript:" not in rendered
    assert '<img src="images/a.png" width="24" height="24"/>' in rendered


def test_v030_defaults():
    cfg = load_config(None)
    assert cfg["confluence"]["page_width"] == "max"
    assert cfg["confluence"]["page_selector"]["max_depth"] == 0
    assert cfg["profiles"]["publication"]["docx"]["suppress_floating_textboxes_in_flow"] is True


def test_docx_toc_marker_and_floating_textbox_suppression():
    from lxml import etree, html
    from docspecbridge.docx_publication import (
        NS,
        TOC_MARKER,
        W,
        _has_toc_field,
        _replace_source_toc,
        _strip_floating_textbox_flow,
    )

    root = etree.fromstring(
        f'''<w:document xmlns:w="{NS['w']}"><w:body>
        <w:p><w:r><w:instrText> TOC \\o "1-3" </w:instrText></w:r></w:p>
        <w:p><w:r><w:drawing><w:txbxContent><w:p><w:r><w:t>2.3</w:t></w:r></w:p></w:txbxContent></w:drawing></w:r></w:p>
        </w:body></w:document>'''.encode()
    )
    assert _has_toc_field(root) is True
    removed = _strip_floating_textbox_flow(root)
    assert removed == ["2.3"]
    assert "2.3" not in "".join(root.xpath(".//w:t/text()", namespaces=NS))

    html_root = html.fragment_fromstring(f"<p>Table des matières</p><p>{TOC_MARKER}</p><h1>1. Objet</h1>", create_parent="div")
    replaced, _ = _replace_source_toc(html_root, "auto", True)
    assert replaced is True
    assert html_root.xpath(".//docspecbridge-toc")


def test_v040_canonical_merged_cells_are_not_duplicated():
    from docspecbridge.canonical import canonical_from_xhtml
    from docspecbridge.renderers import render_markdown, render_rag, render_confluence

    doc = canonical_from_xhtml(
        '<table><tr><td>Version</td><td colspan="3">4.0</td></tr>'
        '<tr><td rowspan="2">Confidentialité</td><td>A</td><td>B</td><td>C</td></tr>'
        '<tr><td></td><td>OUI</td><td>NON</td></tr></table>',
        title="test",
    )
    md = render_markdown(doc)
    rag = render_rag(doc, {})
    confluence = render_confluence(doc)
    assert md.count("4.0") == 1
    assert rag.count("4.0") == 1
    assert 'colspan="3"' in confluence
    assert 'rowspan="2"' in confluence


def test_v040_human_markdown_has_no_table_html():
    from docspecbridge.canonical import canonical_from_xhtml
    from docspecbridge.renderers import render_markdown

    doc = canonical_from_xhtml('<table><tr><td>A</td><td colspan="2">B</td></tr></table>', title="x")
    md = render_markdown(doc)
    assert "<table" not in md
    assert "<td" not in md
    assert md.count("B") == 1


def test_v040_html_destination():
    from docspecbridge.canonical import canonical_from_xhtml
    from docspecbridge.renderers import render_html

    doc = canonical_from_xhtml('<h1>Title</h1><p><strong>Hello</strong></p>', title="x")
    output = render_html(doc)
    assert "<!doctype html>" in output.lower()
    assert "<h1>Title</h1>" in output
    assert "<strong>Hello</strong>" in output


def test_v040_markdown_source_round_trip():
    from docspecbridge.html_io import canonical_from_markdown
    from docspecbridge.renderers import render_markdown

    source = "# Title\n\n**Bold with literal \\***\n\n| A | B |\n|---|---|\n| 1 | 2 |\n"
    doc = canonical_from_markdown(source, title="x")
    output = render_markdown(doc)
    assert "# Title" in output
    assert "| A | B |" in output
    assert "<table" not in output


def test_v040_adf_round_trip_basic():
    from docspecbridge.adf import canonical_from_adf, canonical_to_adf

    adf = {"version": 1, "type": "doc", "content": [
        {"type": "heading", "attrs": {"level": 2}, "content": [{"type": "text", "text": "Scope"}]},
        {"type": "paragraph", "content": [{"type": "text", "text": "Important", "marks": [{"type": "strong"}]}]},
    ]}
    doc = canonical_from_adf(adf, title="X")
    rebuilt = canonical_to_adf(doc)
    assert rebuilt["type"] == "doc"
    assert rebuilt["content"][0]["type"] == "heading"
    assert rebuilt["content"][1]["content"][0]["marks"][0]["type"] == "strong"


def test_v040_human_markdown_toc_is_html_free():
    from docspecbridge.canonical import canonical_from_xhtml
    from docspecbridge.renderers import render_markdown

    doc = canonical_from_xhtml('[[_TOC_]]<h1>Title</h1><p>Text</p>', title='x')
    md = render_markdown(doc)
    assert '<!--' not in md
    assert '<' not in md
    assert '# Title' in md


def test_v040_html_source_drops_script_and_style_payloads():
    from docspecbridge.canonical import canonical_from_html_document
    from docspecbridge.renderers import render_markdown

    doc = canonical_from_html_document(
        '<html><head><style>.secret{color:red}</style><script>alert("x")</script></head>'
        '<body><main><h1>Visible</h1><p>Content</p><script>steal()</script></main></body></html>',
        title='x',
    )
    md = render_markdown(doc)
    assert 'Visible' in md and 'Content' in md
    assert 'alert' not in md and 'steal' not in md and '.secret' not in md


def test_v040_image_display_geometry_keeps_rag_original_and_human_preview_size(tmp_path):
    from PIL import Image
    from docspecbridge.package_io import write_canonical_package

    package = tmp_path / "pkg"
    images = package / "images"
    images.mkdir(parents=True)
    Image.new("RGB", (400, 400), "white").save(images / "icon.png")
    doc = {
        "schema_version": "1.0",
        "title": "Images",
        "source": {"type": "pptx"},
        "blocks": [{"type": "image", "src": "images/icon.png", "alt": "Icon"}],
        "assets": [{
            "file": "images/icon.png",
            "role": "body",
            "display_occurrences": [{"role": "body", "display": {"width_px": 24, "height_px": 24}}],
            "native": {"width_px": 400, "height_px": 400},
        }],
        "diagnostics": {"warnings": []},
    }
    write_canonical_package(
        doc, package,
        publication_profile={"preserve_image_display_size": True, "display_image_directory": "publication_images"},
        rag_profile={"enabled": True, "keep_image_references": True, "chunking": {"enabled": False}},
    )
    human = (package / "Images.md").read_text(encoding="utf-8")
    rag = (package / "Images.rag.md").read_text(encoding="utf-8")
    html = (package / "Images.html").read_text(encoding="utf-8")
    conf = (package / "Images.confluence.md").read_text(encoding="utf-8")
    assert "publication_images/icon__w24_h24.png" in human
    assert "images/icon.png" in rag
    assert 'width="24"' in html and 'height="24"' in html
    assert 'width="24"' in conf and 'height="24"' in conf


def test_v041_render_document_is_at_package_root(tmp_path):
    from docspecbridge.package_io import write_canonical_package

    doc = {
        "schema_version": "1.0",
        "title": "sample",
        "source": {"type": "test"},
        "assets": [],
        "blocks": [
            {"type": "image", "src": "images/sample.png", "alt": "sample", "width": 12, "height": 12}
        ],
    }
    (tmp_path / "images").mkdir()
    (tmp_path / "images" / "sample.png").write_bytes(b"x")
    outputs = write_canonical_package(doc, tmp_path, rag_profile={"enabled": False})
    render = outputs["confluence_markdown"]
    assert render == tmp_path / "sample.confluence.md"
    text = render.read_text(encoding="utf-8")
    assert 'src="images/sample.png"' in text
    assert '../images/' not in text


def test_v041_space_root_uses_real_homepage_id(monkeypatch):
    import pytest
    pytest.importorskip("md2conf")
    import docspecbridge.confluence as confluence

    def fake_get_json(instance, path):
        assert path == "/spaces/42"
        return {"id": "42", "key": "TT", "name": "Test Space", "homepageId": "123456"}

    def fake_paged_get(instance, path, params):
        if path == "/pages/123456/direct-children":
            return []
        raise AssertionError(path)

    monkeypatch.setattr(confluence, "get_confluence_instance", lambda config, name: ("x", {}))
    monkeypatch.setattr(confluence, "_get_json", fake_get_json)
    monkeypatch.setattr(confluence, "_paged_get", fake_paged_get)
    pages = confluence.list_root_pages({"confluence": {"page_selector": {"max_depth": 0}}}, "x", "42")
    assert pages[0]["id"] == "123456"
    assert pages[0]["space_root"] is True
    assert "SPACE-ROOT" not in pages[0]["tree_label"]


def test_v041_selector_numbering_width_is_stable():
    # Regression guard for the formatting rule used by the interactive selector.
    total = 120
    width = len(str(total))
    assert f"[{9:>{width}}]" == "[  9]"
    assert f"[{10:>{width}}]" == "[ 10]"
    assert f"[{120:>{width}}]" == "[120]"


def test_v042_publication_defaults():
    from docspecbridge.config import load_config

    cfg = load_config(None)
    publication = cfg["confluence"]["publication"]
    assert publication["default_mode"] == "replace"
    assert publication["page_title_source"] == "document_title"
    assert publication["add_title_suffix"] == " ({n})"
    assert publication["verify_after_publish"] is True
    assert cfg["confluence"]["write_page_id_to_markdown"] is False


def test_v042_docx_title_detection_prefers_core_title(tmp_path):
    from docx import Document
    from docspecbridge.title_detection import detect_document_title

    path = tmp_path / "technical-name.docx"
    document = Document()
    document.core_properties.title = "Readable specification title"
    document.add_heading("Section one", level=1)
    document.save(path)

    title, source = detect_document_title(path, {"blocks": []})
    assert title == "Readable specification title"
    assert source == "core_properties"


def test_v042_docx_title_detection_can_use_visual_title(tmp_path):
    from docx import Document
    from docx.shared import Pt
    from docspecbridge.title_detection import detect_document_title

    path = tmp_path / "technical-name.docx"
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Readable visual title")
    run.bold = True
    run.font.size = Pt(24)
    document.add_heading("SUIVI", level=1)
    document.save(path)

    title, source = detect_document_title(path, {"blocks": []})
    assert title == "Readable visual title"
    assert source == "document_visual_title"


def test_v042_html_title_prefers_visible_h1():
    from docspecbridge.canonical import canonical_from_html_document

    doc = canonical_from_html_document(
        "<html><head><title>Browser title</title></head>"
        "<body><main><h1>Visible document title</h1><p>Body</p></main></body></html>"
    )
    assert doc["title"] == "Visible document title"
    assert doc["metadata"]["title_source"] == "document_heading"


def _import_confluence_helpers_without_md2conf():
    """Import confluence helpers even in the lightweight test environment."""
    import importlib
    import sys
    import types

    try:
        return importlib.import_module("docspecbridge.confluence")
    except ModuleNotFoundError as exc:
        if exc.name != "md2conf":
            raise

    names = [
        "md2conf", "md2conf.api", "md2conf.environment", "md2conf.options",
        "md2conf.options_converter", "md2conf.publisher",
    ]
    for name in names:
        sys.modules.setdefault(name, types.ModuleType(name))

    class Dummy:
        def __init__(self, *args, **kwargs):
            pass

    sys.modules["md2conf.api"].ConfluenceAPI = Dummy
    sys.modules["md2conf.environment"].ConnectionProperties = Dummy
    sys.modules["md2conf.options"].ConfluencePageID = Dummy
    sys.modules["md2conf.options"].ProcessorOptions = Dummy
    sys.modules["md2conf.options_converter"].ConverterOptions = Dummy
    sys.modules["md2conf.options_converter"].ImageLayoutOptions = Dummy
    sys.modules["md2conf.options_converter"].LayoutOptions = Dummy
    sys.modules["md2conf.options_converter"].TableLayoutOptions = Dummy
    sys.modules["md2conf.publisher"].Publisher = Dummy
    sys.modules.pop("docspecbridge.confluence", None)
    return importlib.import_module("docspecbridge.confluence")


def test_v042_render_document_stays_target_agnostic(tmp_path):
    confluence = _import_confluence_helpers_without_md2conf()
    md = tmp_path / "render_document.md"
    original = (
        '---\ntitle: "Old title"\npage_id: "111"\n---\n'
        '<!-- confluence-page-id: 111 -->\n'
        '<!-- confluence-space-key: OLD -->\n'
        '<p><img src="images/a.png"/></p>\n'
    )
    md.write_text(original, encoding="utf-8")

    routed = confluence._publication_text(md, title="New title", page_id="222")
    # The temporary publication payload gets the requested route.
    assert "page_id: '222'" in routed or 'page_id: "222"' in routed or "page_id: 222" in routed
    assert "title: New title" in routed or "title: 'New title'" in routed or 'title: "New title"' in routed
    assert "confluence-page-id" not in routed
    assert "confluence-space-key" not in routed
    assert 'src="images/a.png"' in routed
    # The generated package renderer itself must never be mutated.
    assert md.read_text(encoding="utf-8") == original


def test_v042_publication_state_separates_primary_and_add_copy(tmp_path):
    confluence = _import_confluence_helpers_without_md2conf()
    md = tmp_path / "render_document.md"
    md.write_text("---\ntitle: X\n---\n", encoding="utf-8")

    confluence._save_publication_state(
        md, instance_name="prod", space_key="DOC", space_id="1", parent_id="10",
        page_id="100", title="X", publication_mode="replace",
    )
    confluence._save_publication_state(
        md, instance_name="prod", space_key="DOC", space_id="1", parent_id="10",
        page_id="101", title="X (2)", publication_mode="add",
    )
    primary = confluence._state_entry(
        md, instance_name="prod", space_key="DOC", parent_id="10"
    )
    state = confluence._load_publication_state(md)
    assert primary["page_id"] == "100"
    assert {item["role"] for item in state["confluence"]} == {"primary", "copy"}
    assert len(state["confluence"]) == 2



def test_v050_source_named_outputs_and_legacy_resolution(tmp_path):
    import json
    from docspecbridge.package_io import write_canonical_package, resolve_package_output

    doc = {"schema_version": "1.0", "title": "Specification", "source": {"type": "test"}, "blocks": [], "assets": [], "diagnostics": {"warnings": []}}
    outputs = write_canonical_package(doc, tmp_path, stem="specification", rag_profile={"enabled": True, "chunking": {"enabled": False}})
    assert outputs["human_markdown"].name == "specification.md"
    assert outputs["document_json"].name == "specification.json"
    assert outputs["html"].name == "specification.html"
    assert outputs["rag_markdown"].name == "specification.rag.md"
    assert outputs["confluence_markdown"].name == "specification.confluence.md"
    manifest = json.loads((tmp_path / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["schema_version"] == "0.5"

    legacy = tmp_path / "legacy"
    legacy.mkdir()
    (legacy / "document.json").write_text("{}", encoding="utf-8")
    assert resolve_package_output(legacy, "document_json", "document.json") == legacy / "document.json"


def test_v050_adf_external_media_keeps_image_position():
    from docspecbridge.adf import canonical_to_adf
    doc = {
        "blocks": [
            {"type": "paragraph", "inlines": [{"type": "text", "text": "before", "marks": []}]},
            {"type": "image", "src": "images/a.png", "alt": "A", "width": 320, "height": 200},
            {"type": "paragraph", "inlines": [{"type": "text", "text": "after", "marks": []}]},
        ]
    }
    adf = canonical_to_adf(doc, {"images/a.png": "https://example.atlassian.net/rest/api/3/attachment/content/1"})
    assert [node["type"] for node in adf["content"]] == ["paragraph", "mediaSingle", "paragraph"]
    media = adf["content"][1]["content"][0]["attrs"]
    assert media["type"] == "external"
    assert media["width"] == 320


def test_v050_xlsx_extracts_sheets_formulas_and_chart_preview(tmp_path):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from docspecbridge.config import load_config
    from docspecbridge.xlsx_io import extract_xlsx_to_package

    source = tmp_path / "budget.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Month", "Amount"])
    ws.append(["Jan", 10])
    ws.append(["Feb", 20])
    ws["B4"] = "=SUM(B2:B3)"
    chart = BarChart()
    chart.title = "Amounts"
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=3))
    ws.add_chart(chart, "D2")
    wb.create_sheet("Notes")["A1"] = "hello"
    wb.save(source)

    package = tmp_path / "out" / "budget__xlsx"
    package.mkdir(parents=True)
    result = extract_xlsx_to_package(source, package, load_config(None))
    assert (package / "budget.md").is_file()
    assert len(result["children"]) == 2
    first = package / result["children"][0]["package"]
    assert (first / "Budget.md").is_file()
    assert list((first / "images").glob("chart_*.svg"))
    meta = result["metadata"]["sheets"][0]
    assert any(x["cell"] == "B4" and x["formula"].startswith("=") for x in meta["formulas"])


def test_v050_config_has_jira_resume_and_confluence_attachment_modes():
    from docspecbridge.config import load_config
    cfg = load_config(None)
    assert cfg["jira"]["import"]["resume"] is True
    assert cfg["jira"]["export"]["include_comments"] is True
    assert cfg["confluence"]["publication"]["attachments"] == "none"
    assert cfg["confluence"]["export"]["attachment_mode"] == "all"


def test_cli_version_option():
    from typer.testing import CliRunner
    from docspecbridge.cli import app

    result = CliRunner().invoke(app, ["--version"])
    assert result.exit_code == 0
    assert result.stdout.strip() == "DocSpecBridge 0.5.1"

    short = CliRunner().invoke(app, ["-V"])
    assert short.exit_code == 0
    assert short.stdout.strip() == "DocSpecBridge 0.5.1"


def test_v050_all_new_interactive_menu_labels_are_localized():
    from docspecbridge.i18n import SUPPORTED_LANGUAGES, tr

    keys = [
        "main.extract_v050", "main.import_v050", "extract.title", "extract.local",
        "extract.confluence", "extract.jira", "import.title", "import.confluence",
        "import.jira", "rag.export_existing", "interactive.space_select",
        "interactive.page_select", "interactive.project_select", "interactive.issue_type_select",
        "jira.settings", "jira.settings.default_project", "jira.settings.default_issue_type",
        "interactive.cancelled",
    ]
    for lang in SUPPORTED_LANGUAGES:
        cfg = load_config(None)
        cfg["app"]["language"] = lang
        for key in keys:
            value = tr(cfg, key)
            assert value != key
            assert value.strip()


def test_v050_jira_instances_merge_confluence_and_jira_overrides():
    from docspecbridge.jira import jira_instances

    cfg = load_config(None)
    cfg["confluence"]["instances"] = {
        "prod": {"domain": "prod.atlassian.net", "user_name": "a@example.com"},
        "sandbox": {"domain": "sandbox.atlassian.net", "user_name": "a@example.com"},
    }
    cfg["jira"]["instances"] = {
        "prod": {"domain": "jira-prod.atlassian.net", "user_name": "b@example.com"},
        "jira-only": {"domain": "jira-only.atlassian.net", "user_name": "b@example.com"},
    }
    instances = jira_instances(cfg)
    assert set(instances) == {"prod", "sandbox", "jira-only"}
    assert instances["prod"]["domain"] == "jira-prod.atlassian.net"
    assert instances["sandbox"]["domain"] == "sandbox.atlassian.net"


def test_v050_default_config_has_jira_interactive_defaults():
    cfg = load_config(None)
    assert "defaults" in cfg["jira"]
    assert cfg["jira"]["defaults"] == {}


def test_v050_confluence_location_discovery_keeps_defaults_visible(monkeypatch):
    import docspecbridge.cli as cli

    cfg = load_config(None)
    cfg["confluence"]["instances"] = {
        "prod": {"default_space": "B", "root_page": "22"}
    }
    cfg["confluence"]["default_instance"] = "prod"
    monkeypatch.setattr(cli, "list_spaces", lambda cfg, instance: [
        {"id": "1", "key": "A", "name": "Alpha", "homepageId": "11"},
        {"id": "2", "key": "B", "name": "Beta", "homepageId": "21"},
    ])
    monkeypatch.setattr(cli, "list_root_pages", lambda cfg, instance, sid, max_depth=0: [
        {"id": "21", "title": "Home", "tree_label": "Home"},
        {"id": "22", "title": "Root", "tree_label": "Root"},
    ])
    seen = []

    def fake_select(title, options, default_index=0, **kwargs):
        opts = list(options)
        seen.append((title, default_index, opts))
        return opts[default_index][0]

    monkeypatch.setattr(cli, "select_option", fake_select)
    space, page, _ = cli._select_confluence_location(cfg, "prod")
    assert space["key"] == "B"
    assert page["id"] == "22"
    assert seen[0][1] == 1  # default space B is highlighted, not bypassed
    assert seen[1][1] == 1  # default root page 22 is highlighted, not bypassed


def test_v050_jira_project_discovery_uses_configured_default(monkeypatch):
    import docspecbridge.cli as cli

    cfg = load_config(None)
    cfg["jira"]["defaults"] = {"prod": {"project": "B"}}
    monkeypatch.setattr(cli, "list_projects", lambda cfg, instance, query=None: [
        {"id": "1", "key": "A", "name": "Alpha"},
        {"id": "2", "key": "B", "name": "Beta"},
    ])
    seen = {}

    def fake_select(title, options, default_index=0, **kwargs):
        opts = list(options)
        seen["default_index"] = default_index
        return opts[default_index][0]

    monkeypatch.setattr(cli, "select_option", fake_select)
    project = cli._select_jira_project(cfg, "prod")
    assert project["key"] == "B"
    assert seen["default_index"] == 1


def test_v050_interactive_prompts_use_escape_aware_ui_layer():
    from pathlib import Path
    root = Path(__file__).parent / "src" / "docspecbridge"
    assert "Prompt.ask" not in (root / "cli.py").read_text(encoding="utf-8")
    assert "Prompt.ask" not in (root / "config_ui.py").read_text(encoding="utf-8")


def test_v050_jira_discovery_default_page_size():
    cfg = load_config(None)
    assert cfg["jira"]["discovery"]["page_size"] == 50


def test_v050_jira_discovery_labels_are_localized():
    from docspecbridge.i18n import SUPPORTED_LANGUAGES, tr
    keys = [
        "jira.discovery.title", "jira.discovery.search", "jira.discovery.reset",
        "jira.discovery.manual", "jira.discovery.next", "jira.discovery.search_prompt",
        "jira.discovery.manual_prompt", "jira.discovery.mode_title",
        "jira.discovery.mode_manual", "jira.discovery.mode_browse", "jira.discovery.title_typed",
    ]
    for lang in SUPPORTED_LANGUAGES:
        cfg = load_config(None)
        cfg["app"]["language"] = lang
        for key in keys:
            if key == "jira.discovery.title":
                assert tr(cfg, key, project="ABC", count=2) != key
            elif key == "jira.discovery.title_typed":
                assert tr(cfg, key, project="ABC", issue_type="Story", count=2) != key
            else:
                assert tr(cfg, key) != key


def test_v050_jira_issue_selector_lists_recent_issues_without_loading_whole_project(monkeypatch):
    import docspecbridge.cli as cli

    cfg = load_config(None)
    calls = []
    monkeypatch.setattr(cli, "list_issues", lambda cfg, project, instance, issue_type=None, query=None, next_page_token=None, max_results=50: (
        calls.append((project, issue_type, query, next_page_token, max_results)) or {
            "issues": [{"key": "ABC-2", "fields": {"summary": "Second", "status": {"name": "Open"}, "issuetype": {"name": "Story"}, "updated": "2026-09-20T12:00:00.000+0000"}}],
            "next_page_token": "NEXT",
            "is_last": False,
        }
    ))
    monkeypatch.setattr(cli, "select_option", lambda title, options, **kwargs: list(options)[0][0])
    selected = cli._select_jira_issue(cfg, "prod", "ABC", "Story")
    assert selected == "ABC-2"
    assert calls == [("ABC", "Story", None, None, 50)]


def test_v050_jira_media_rewrite_matches_adf_alt_filename_and_links():
    from docspecbridge.jira import _rewrite_downloaded_media
    doc = {
        "blocks": [
            {"type": "image", "src": "media-uuid", "alt": "shot.png", "media": {"id": "media-uuid"}},
            {"type": "paragraph", "inlines": [{"type": "text", "text": "spec.pdf", "marks": [{"type": "link", "href": "https://example/attachment/spec.pdf"}]}]},
        ]
    }
    attachments = [
        {"saved": True, "id": "100", "filename": "shot.png", "file": "attachments/shot.png"},
        {"saved": True, "id": "101", "filename": "spec.pdf", "file": "attachments/spec.pdf"},
    ]
    _rewrite_downloaded_media(doc, attachments)
    assert doc["blocks"][0]["src"] == "attachments/shot.png"
    assert doc["blocks"][1]["inlines"][0]["marks"][0]["href"] == "attachments/spec.pdf"


def test_v050_jira_canonical_contains_comments_and_visible_attachment_section():
    from docspecbridge.jira import issue_to_canonical
    issue = {"key": "ABC-1", "fields": {"summary": "Demo", "description": {"type": "doc", "version": 1, "content": []}}}
    doc = issue_to_canonical(
        issue,
        instance_name="prod",
        comments=[{"author": {"displayName": "Alice"}, "created": "2026-09-20", "body": "Plain comment"}],
        downloaded_attachments=[
            {"saved": True, "filename": "shot.png", "mime_type": "image/png", "file": "attachments/shot.png"},
            {"saved": True, "filename": "spec.pdf", "mime_type": "application/pdf", "file": "attachments/spec.pdf"},
        ],
        comments_title="Commentaires",
        attachments_title="Pièces jointes",
    )
    headings = ["".join(i.get("text", "") for i in b.get("inlines", [])) for b in doc["blocks"] if b.get("type") == "heading"]
    assert "Commentaires" in headings
    assert "Pièces jointes" in headings
    assert any(b.get("type") == "paragraph" and any(i.get("text") == "Plain comment" for i in b.get("inlines", [])) for b in doc["blocks"])
    assert any(b.get("type") == "image" and b.get("src") == "attachments/shot.png" for b in doc["blocks"])
    assert any(any(m.get("href") == "attachments/spec.pdf" for i in b.get("inlines", []) for m in i.get("marks", [])) for b in doc["blocks"] if b.get("type") == "paragraph")


def test_v050_jira_attachment_download_uses_dedicated_content_endpoint(monkeypatch, tmp_path):
    import docspecbridge.jira as jira

    seen = {}
    class Response:
        content = b"payload"
        def raise_for_status(self): pass
    class Client:
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def get(self, url, params=None):
            seen["url"] = url
            seen["params"] = params
            return Response()
    monkeypatch.setattr(jira, "_client", lambda instance, timeout=120.0: Client())
    issue = {"fields": {"attachment": [{"id": "123", "filename": "a.txt", "mimeType": "text/plain", "content": "https://remote.example/file"}]}}
    assets = jira._download_attachments({"domain": "example.atlassian.net"}, issue, tmp_path)
    assert seen["url"].endswith("/rest/api/3/attachment/content/123")
    assert seen["params"] == {"redirect": "false"}
    assert (tmp_path / "attachments" / "a.txt").read_bytes() == b"payload"
    assert assets[0]["saved"] is True


def test_v050_jira_extract_manual_mode_does_not_discover_project(monkeypatch, tmp_path):
    import docspecbridge.cli as cli
    cfg = load_config(None)
    cfg["app"]["destination"] = str(tmp_path)
    monkeypatch.setattr(cli, "_select_jira_instance", lambda cfg: "prod")
    answers = iter(["manual"])
    monkeypatch.setattr(cli, "select_option", lambda *a, **k: next(answers))
    monkeypatch.setattr(cli, "prompt_text", lambda *a, **k: "ABC-123")
    monkeypatch.setattr(cli, "_select_jira_project", lambda *a, **k: (_ for _ in ()).throw(AssertionError("project discovery must not run")))
    seen = {}
    monkeypatch.setattr(cli, "export_issue", lambda cfg, issue, destination, instance_name=None: seen.update(issue=issue, instance=instance_name) or destination / "pkg")
    cli._interactive_extract_jira(cfg)
    assert seen == {"issue": "ABC-123", "instance": "prod"}


def test_v050_jira_extract_browse_selects_type_before_issue_list(monkeypatch, tmp_path):
    import docspecbridge.cli as cli
    cfg = load_config(None)
    cfg["app"]["destination"] = str(tmp_path)
    monkeypatch.setattr(cli, "_select_jira_instance", lambda cfg: "prod")
    monkeypatch.setattr(cli, "select_option", lambda *a, **k: "browse")
    order = []
    monkeypatch.setattr(cli, "_select_jira_project", lambda *a, **k: order.append("project") or {"key": "ABC"})
    monkeypatch.setattr(cli, "_select_jira_issue_type", lambda *a, **k: order.append("type") or "Story")
    monkeypatch.setattr(cli, "_select_jira_issue", lambda cfg, instance, project, issue_type: order.append(("issues", project, issue_type)) or "ABC-2")
    monkeypatch.setattr(cli, "export_issue", lambda cfg, issue, destination, instance_name=None: destination / "pkg")
    cli._interactive_extract_jira(cfg)
    assert order == ["project", "type", ("issues", "ABC", "Story")]


def test_v050_jira_list_issues_accepts_issue_type_filter_in_source():
    from pathlib import Path
    source = (Path(__file__).parent / "src" / "docspecbridge" / "jira.py").read_text(encoding="utf-8")
    assert 'issuetype = "{escaped_type}"' in source


def test_v051_legacy_extensions_are_removed_from_config(tmp_path):
    import yaml
    from docspecbridge.config import load_config
    cfg_path = tmp_path / "docspecbridge.yaml"
    cfg_path.write_text(yaml.safe_dump({"app": {"extensions": [".docx", ".pdf"]}}), encoding="utf-8")
    cfg = load_config(cfg_path)
    assert "extensions" not in cfg["app"]


def test_v051_save_does_not_persist_extensions(tmp_path):
    import yaml
    from docspecbridge.config import load_config, save_config
    cfg_path = tmp_path / "docspecbridge.yaml"
    cfg_path.write_text(yaml.safe_dump({"app": {"extensions": [".docx"]}}), encoding="utf-8")
    cfg = load_config(cfg_path)
    save_config(cfg, cfg_path)
    saved = yaml.safe_load(cfg_path.read_text(encoding="utf-8"))
    assert "extensions" not in saved["app"]


def test_v051_xlsx_is_discovered_even_with_legacy_yaml(tmp_path):
    import yaml
    from docspecbridge.config import load_config
    from docspecbridge.extractor import XbergExtractor
    source = tmp_path / "input"
    source.mkdir()
    (source / "book.xlsx").write_bytes(b"placeholder")
    cfg_path = tmp_path / "docspecbridge.yaml"
    cfg_path.write_text(yaml.safe_dump({"app": {"extensions": [".docx", ".pdf"]}}), encoding="utf-8")
    cfg = load_config(cfg_path)
    files, _ = XbergExtractor(cfg).discover(source)
    assert [p.name for p in files] == ["book.xlsx"]


def test_v051_cli_extension_filter_is_runtime_only(tmp_path):
    from docspecbridge.config import load_config
    from docspecbridge.cli import _runtime_config
    from docspecbridge.extractor import XbergExtractor
    source = tmp_path / "input"
    source.mkdir()
    (source / "a.pdf").write_bytes(b"pdf")
    (source / "b.xlsx").write_bytes(b"xlsx")
    cfg = _runtime_config(load_config(None), extensions=["pdf"])
    assert "extensions" not in cfg["app"]
    files, _ = XbergExtractor(cfg).discover(source)
    assert [p.name for p in files] == ["a.pdf"]


def test_v050_jira_comments_merge_platform_embedded_and_jsm(monkeypatch):
    import docspecbridge.jira as jira
    cfg = {"jira": {"instances": {"prod": {"domain": "example.atlassian.net", "user_name": "u"}}, "default_instance": "prod"}}
    issue = {
        "fields": {
            "project": {"projectTypeKey": "service_desk"},
            "comment": {"total": 2, "comments": [{"id": "1", "body": "embedded", "author": {"displayName": "A"}}]},
        }
    }
    monkeypatch.setattr(jira, "get_comments", lambda *a, **k: [{"id": "1", "body": "platform", "author": {"displayName": "A"}}])
    monkeypatch.setattr(jira, "_jsm_comments", lambda *a, **k: [{"id": "2", "body": "jsm", "author": {"displayName": "B"}}])
    comments, warnings = jira.collect_comments(cfg, "ABC-1", issue, "prod")
    assert [c["id"] for c in comments] == ["1", "2"]
    assert warnings == []


def test_v050_jira_comments_warn_when_issue_reports_comments_but_api_returns_none(monkeypatch):
    import docspecbridge.jira as jira
    cfg = {"jira": {"instances": {"prod": {"domain": "example.atlassian.net", "user_name": "u"}}, "default_instance": "prod"}}
    issue = {"fields": {"project": {"projectTypeKey": "software"}, "comment": {"total": 3, "comments": []}}}
    monkeypatch.setattr(jira, "get_comments", lambda *a, **k: [])
    monkeypatch.setattr(jira, "_jsm_comments", lambda *a, **k: [])
    comments, warnings = jira.collect_comments(cfg, "ABC-1", issue, "prod")
    assert comments == []
    assert any("reports 3 comment" in warning for warning in warnings)
